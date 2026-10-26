"""G1 workbook import (FN-16 seed path): template <-> parser round trip and error reporting."""

import io
import shutil
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.quote_engine.g1 import FIRST_ROW, SHEETS, G1Error, build_template, parse_g1

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs" / "inputs" / "G1_입력양식.xlsx"

FILLED = {
    "1_재질": [
        ("SS400", "일반구조용 압연강", 7.85, 1200.5, 5, "전체", None),
        ("SUS304", None, 7.93, 4500, "3%", "1.0~3.2", None),
    ],
    "3_임률": [("레이저", 35000, None), ("절곡", 30000, None)],
    "2_공정": [
        (
            "LASER",
            "레이저 절단",
            "cutting_length_mm",
            "setup_min + cutting_length_mm / speed",
            "setup_min=5, speed=3000",
            60000,
            "레이저",
            None,
        ),
        (
            "BEND",
            "절곡",
            "bend_count",
            "setup_min + bend_count * per_bend",
            "setup_min=3, per_bend=0.5",
            None,
            "절곡",
            None,
        ),
    ],
    "5_레이어": [
        ("절곡선", "BEND*, FOLD", "DASHED", None),
        ("무시", "DIM*", None, None),
        ("펀칭 기준 직경(mm)", "≤ 6", None, None),
    ],
    "6_표제란": [("재질", "MAT_CODE", "재질, MATERIAL", None)],
}
RATES = {
    "제조간접비 배부 기준": "기계시간",
    "일반관리비율(%)": 6,
    "이윤율(%)": 10,
    "절사 규칙": "절사",
    "절사 단위(원)": 10,
    "절사 적용 시점": "합계에서 한 번",
    "재료 중량 기준": "소재 사각",
}


def fill(tmp_path: Path, rows=FILLED, rates=RATES, name="g1.xlsx") -> str:
    out = tmp_path / name
    shutil.copy(TEMPLATE, out)
    wb = load_workbook(out)
    for sheet, data in rows.items():
        ws = wb[sheet]
        start = FIRST_ROW + len(SHEETS[sheet][2])  # first yellow row, after the examples
        for i, row in enumerate(data):
            for c, v in enumerate(row, 1):
                ws.cell(row=start + i, column=c, value=v)
    ws = wb["4_원가비율"]
    for r in range(FIRST_ROW, FIRST_ROW + len(SHEETS["4_원가비율"][2])):
        if (label := ws.cell(row=r, column=1).value) in rates:
            ws.cell(row=r, column=2, value=rates[label])
    wb.save(out)
    return str(out)


def test_committed_template_is_current(tmp_path):
    # docs/inputs/G1_입력양식.xlsx must be regenerated whenever SHEETS changes
    fresh = tmp_path / "fresh.xlsx"
    build_template(str(fresh))
    for title, (_, cols, _) in SHEETS.items():
        a, b = load_workbook(TEMPLATE)[title], load_workbook(fresh)[title]
        assert [c.value for c in a[4]][: len(cols)] == [c.value for c in b[4]][: len(cols)], title


def test_blank_template_imports_nothing():
    r = parse_g1(str(TEMPLATE))
    assert (
        r["errors"] == [] and r["bundle"]["materials"] == [] and r["bundle"]["cost_ratios"] is None
    )
    assert r["warnings"] and "예시" in r["warnings"][0]


def test_filled_workbook_becomes_a_bundle(tmp_path):
    r = parse_g1(fill(tmp_path))
    assert r["errors"] == []
    b = r["bundle"]
    ss, sus = b["materials"]
    assert ss == {
        "material_code": "SS400",
        "material_name": "일반구조용 압연강",
        "thickness_min_mm": None,
        "thickness_max_mm": None,
        "density_g_cm3": "7.85",
        "unit_price_per_kg": "1200.5",
        "scrap_rate": "0.05",
    }
    assert sus["scrap_rate"] == "0.03" and (sus["thickness_min_mm"], sus["thickness_max_mm"]) == (
        "1.0",
        "3.2",
    )
    items = {p["item_code"]: p for p in b["price_items"]}
    assert items["레이저"]["item_type"] == "LABOR" and items["LASER-MC"] == {
        "item_code": "LASER-MC",
        "item_type": "MACHINE",
        "unit": "h",
        "unit_price": "60000",
    }
    laser = b["process_rules"][0]
    assert laser["params"] == {"setup_min": 5.0, "speed": 3000.0}
    assert laser["labor_item_code"] == "레이저" and laser["machine_item_code"] == "LASER-MC"
    assert b["cost_ratios"] == {
        "overhead_basis": "MACHINE_HOUR",
        "admin_rate": "0.06",
        "profit_rate": "0.1",
        "vat_rate": "0.1",
        "rounding_rule": "FLOOR",
        "rounding_unit": 10,
        "rounding_scope": "TOTAL",
        "material_basis": "BBOX",
        "stock_allowance_mm": "2",  # the template's default (P4-1)
    }
    rules = {(m["rule_type"], m["target"], m["pattern"]) for m in b["mapping_rules"]}
    assert rules == {
        ("LAYER", "BEND", "BEND*, FOLD"),
        ("LINETYPE", "BEND", "DASHED"),
        ("LAYER", "IGNORE", "DIM*"),
        ("PUNCH_MAX_DIA", "HOLE", "6"),
        ("TITLE_TAG", "material", "MAT_CODE, 재질, MATERIAL"),
    }


def test_construction_sheets_import(tmp_path):
    rows = FILLED | {
        "9_공종단가": [
            (
                "A-FLOOR",
                "건축",
                "바닥 타일",
                "300각",
                "m2",
                "room_area_m2",
                3,
                20000,
                15000,
                0,
                "2027-12-31",
                None,
            ),
        ],
        "10_보정계수": [("지역", "SEOUL", 1.05, None), ("시점", "2027Q1", 1.02, None)],
        "5_레이어": [*FILLED["5_레이어"], ("벽체", "W-*, 벽*", None, None)],
    }
    r = parse_g1(fill(tmp_path, rows, RATES | {"간접노무비율(%)": 8, "간접공사비율(%)": 5}))
    assert r["errors"] == []
    b = r["bundle"]
    assert b["work_items"] == [
        {
            "item_code": "A-FLOOR",
            "work_category": "건축",
            "item_name": "바닥 타일",
            "spec": "300각",
            "unit": "m2",
            "qty_formula": "room_area_m2",
            "loss_rate": "0.03",
            "material_price": "20000",
            "labor_price": "15000",
            "expense_price": "0",
            "valid_until": "2027-12-31",
            "costbook_item_code": None,
        }
    ]
    assert [(f["factor_type"], f["factor"]) for f in b["adjust_factors"]] == [
        ("REGION", "1.05"),
        ("ESCALATION", "1.02"),
    ]
    assert ("QTO_LAYER", "WALL", "W-*, 벽*") in {
        (m["rule_type"], m["target"], m["pattern"]) for m in b["mapping_rules"]
    }
    assert b["cost_ratios"]["indirect_labor_rate"] == "0.08"


def test_workbook_without_construction_sheets_still_imports(tmp_path):
    out = fill(tmp_path)
    wb = load_workbook(out)
    del wb["9_공종단가"], wb["10_보정계수"]
    wb.save(out)
    r = parse_g1(out)
    assert r["errors"] == [] and r["bundle"]["work_items"] == []


def test_row_errors_are_collected_with_sheet_and_row(tmp_path):
    bad = {
        "1_재질": [("SS400", None, "철", None, None, "전체", None)],
        "2_공정": [("LASER", None, "cutting_length_mm", "x", "speed=?", None, "없는구분", None)],
        "5_레이어": [("외곽", "0", None, None)],
    }
    r = parse_g1(fill(tmp_path, bad, {"절사 규칙": "버림", "재료 중량 기준": "순면적 / 소재 사각"}))
    text = "\n".join(r["errors"])
    assert "1_재질 8행: 숫자가 아닙니다: 철" in text
    assert "2_공정 15행" in text and "5_레이어 10행" in text  # 2_공정: after its example rows
    assert "4_원가비율" in text  # unknown choice
    assert "여러 개" in text  # the hint copied verbatim is ambiguous, not silently NET


def test_rejects_non_workbooks_and_tampered_files(tmp_path):
    junk = tmp_path / "junk.xlsx"
    junk.write_bytes(b"not a zip")
    with pytest.raises(G1Error):
        parse_g1(str(junk))
    bomb = tmp_path / "bomb.xlsx"  # inflates past the 50 MB cap
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/workbook.xml", "<x/>")
        z.writestr("big.bin", b"\0" * (60 * 1024**2))
    with pytest.raises(G1Error, match="너무 큽니다"):
        parse_g1(str(bomb))
    path = fill(tmp_path)
    wb = load_workbook(path)
    wb["1_재질"].cell(row=4, column=1, value="코드")
    del wb["6_표제란"]
    buf = io.BytesIO()
    wb.save(buf)
    (tmp_path / "t.xlsx").write_bytes(buf.getvalue())
    with pytest.raises(G1Error, match="6_표제란"):
        parse_g1(str(tmp_path / "t.xlsx"))


def test_costbook_sheets(tmp_path):
    """P4-6: 표준품셈 items with their source, resources per unit and resource prices; a work item
    linked to an item; a workbook filled in before the 품셈코드 column still imports."""
    rows = {
        "9_공종단가": [
            (
                "A-PLASTER",
                "건축",
                "벽 미장",
                None,
                "m2",
                "wall_length_m * height_m * 2",
                0,
                None,
                None,
                None,
                None,
                None,
                "CB-PLASTER-W",
            ),
        ],
        "12_품셈항목": [
            (
                "CB-PLASTER-W",
                "시멘트 모르타르 미장",
                "벽, 18mm",
                "m2",
                "미장",
                "국토교통부 공고 제2025-1574호",
                "2026-01-01",
                512,
                "A" * 64,
                "수기",
                "홍길동",
                "2026-01-05",
                None,
            ),
        ],
        "13_품셈투입량": [
            ("CB-PLASTER-W", "노무", "L-PLASTERER", 0.06, None),
            ("CB-PLASTER-W", "재료", "M-CEMENT", 8.5, None),
        ],
        "14_자원단가": [
            (
                "L-PLASTERER",
                "노무",
                "미장공",
                "인",
                280000,
                "2026-01-01",
                "시중노임단가",
                None,
                None,
            ),
            ("M-CEMENT", "재료", "시멘트", "kg", 150, None, None, None, None),
        ],
    }
    r = parse_g1(fill(tmp_path, rows))
    assert r["errors"] == []
    b = r["bundle"]
    assert b["work_items"][0]["costbook_item_code"] == "CB-PLASTER-W"
    (item,) = b["costbook_items"]
    assert item["page_no"] == 512 and item["source_sha256"] == "a" * 64
    assert item["extract_method"] == "MANUAL" and item["notice_no"].endswith("제2025-1574호")
    assert [(x["resource_type"], x["qty_per_unit"]) for x in b["costbook_resources"]] == [
        ("LABOR", "0.06"),
        ("MATERIAL", "8.5"),
    ]
    assert b["resource_prices"][0]["unit_price"] == "280000"
    # an older workbook: no 품셈코드 column on 9_공종단가
    old = fill(tmp_path, {}, name="old.xlsx")
    wb = load_workbook(old)
    wb["9_공종단가"].delete_cols(13)
    wb.save(old)
    assert parse_g1(old)["errors"] == []
