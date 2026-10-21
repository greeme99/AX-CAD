"""FN-21 quote documents through the API: TC-79 (PDF text), TC-80 (on demand, Track 7: no
approval step), TC-81 (XLSX cells)."""

import io
import json
from decimal import Decimal as D

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from backend.api import routes_quote
from core.quote_engine.report import korean_amount
from tests.test_api_quote import estimator  # noqa: F401 - fixture
from tests.test_api_quotes import active, upload  # noqa: F401 - fixture

BODY = {"qty": 3, "material_code": "SS400", "thickness_mm": "2"}
WHY = "고객 협의 단가 적용"


def new_quote(client, world, estimator, **body):  # noqa: F811
    rid = upload(client, world)
    r = client.post("/api/quotes", json={"revision_id": rid, **body}, headers=estimator)
    assert r.status_code == 200, r.text
    return r.json()["data"]


REAL = {
    "company": "(주)에이엑스정밀",
    "business_no": "123-45-67890",
    "ceo": "김대표",
    "address": "경기도 안산시 단원구 공단로 1",
    "phone": "031-123-4567",
    "email": "quote@ax.example",
    "validity_days": 30,
    "delivery": "발주 후 14일",
    "payment": "익월 말 현금",
    "note": "운송비 별도",
}


@pytest.fixture
def supplier(tmp_path, monkeypatch):
    f = tmp_path / "supplier.json"
    f.write_text(json.dumps(REAL, ensure_ascii=False), encoding="utf-8")
    monkeypatch.setattr(routes_quote, "SUPPLIER_FILE", f)
    return f


def pdf_text(data):
    return "".join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)


def test_tc80_documents_on_demand(client, world, estimator, active, make_user, headers):  # noqa: F811
    q = new_quote(client, world, estimator, **BODY)
    url = f"/api/quotes/{q['quote_id']}/report"
    r = client.get(url, headers=estimator)  # no approval needed (Track 7)
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.headers["content-disposition"].endswith(f'"{q["quote_no"]}.pdf"')
    assert r.headers["cache-control"] == "no-store"
    body = pdf_text(r.content)
    assert "샘플 공급자" in body and "산출근거" not in body  # sample supplier is marked
    assert "내부용·대외비" in pdf_text(client.get(f"{url}?basis=true", headers=estimator).content)
    # roles and membership
    for role in ("designer", "viewer"):
        assert client.get(url, headers=world.h[role]).status_code == 403
    uid = make_user("maker", "MANUFACTURING")  # reads quotes, does not export them
    client.post(
        f"/api/projects/{world.pid}/members", json={"user_id": uid}, headers=world.h["admin"]
    )
    assert client.get(url, headers=headers("maker")).status_code == 403
    make_user("outsider", "ESTIMATOR")
    assert client.get(url, headers=headers("outsider")).status_code == 404
    assert client.get(f"{url}?format=docx", headers=estimator).status_code == 422
    assert client.get(f"{url}?official=true", headers=estimator).status_code == 200  # ignored


def test_tc79_81_documents(client, world, estimator, active, supplier):  # noqa: F811
    q = new_quote(client, world, estimator, **BODY)
    line = q["lines"][0]["quote_line_id"]
    adj = {"field": "amount", "value": "9000", "reason": WHY}
    q = client.patch(f"/api/quote-lines/{line}", json=adj, headers=estimator).json()["data"]
    eff = q["effective"]
    url = f"/api/quotes/{q['quote_id']}/report"

    r = client.get(url, headers=world.h["reviewer"])
    assert r.status_code == 200, r.text
    body = pdf_text(r.content)
    total = D(eff["total_amount"])
    assert "샘플" not in body and REAL["company"] in body and REAL["business_no"] in body
    assert f"{D(eff['supply_amount']):,.0f}" in body and f"{total:,.0f}" in body
    assert f"일금 {korean_amount(int(total))}원정" in body
    assert "산출근거" not in body and WHY not in body  # internal basis only on request

    x = client.get(f"{url}?format=xlsx", headers=estimator)
    assert x.headers["content-disposition"].endswith(f'"{q["quote_no"]}.xlsx"')
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["견적서"]
    item = next(r for r in wb["견적서"].iter_rows(values_only=True) if r[0] == 1)
    assert item[3] == BODY["qty"] and D(str(item[6])) == D(eff["supply_amount"])
    assert D(str(item[7])) == D(eff["vat_amount"])

    # always the quote's current values: a later adjustment shows in the next document
    adj["value"] = "19000"
    q = client.patch(f"/api/quote-lines/{line}", json=adj, headers=estimator).json()["data"]
    again = pdf_text(client.get(url, headers=estimator).content)
    assert f"{D(q['effective']['total_amount']):,.0f}" in again
