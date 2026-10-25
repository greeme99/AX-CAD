"""Track 6 construction quotes over the API: QTO preview with a base revision (demolition/new/keep),
construction quote lines and traces, adjust factors, validity warnings, scenarios and sensitivity."""

import io
from decimal import Decimal as D

import pytest
from pypdf import PdfReader

from tests.test_api_master import CONSTRUCTION, FULL, WORK
from tests.test_api_quote import estimator  # noqa: F401 - fixture
from tests.test_api_quote_report import supplier  # noqa: F401 - fixture
from tests.test_qto import W1, W2, W3, plan


def dxf(tmp_path, name, walls, doors=1):
    return io.BytesIO(open(plan(tmp_path / name, walls, doors), "rb").read())


def upload(client, w, data):
    r = client.post(
        f"/api/documents/{w.doc}/dxf", files={"file": ("p.dxf", data)}, headers=w.h["designer"]
    )
    assert r.status_code == 200, r.text
    return r.json()["data"]["revision_id"]


@pytest.fixture
def revs(client, world, tmp_path):
    """(base, current): the existing building (W1, W2, one door) and the renovation (W1, W3,
    two doors) as two revisions of the same drawing."""
    base = upload(client, world, dxf(tmp_path, "base.dxf", [W1, W2]))
    cur = upload(client, world, dxf(tmp_path, "cur.dxf", [W1, W3], doors=2))
    return base, cur


def test_qto_preview_splits_against_the_base(client, world, estimator, revs):  # noqa: F811
    base, cur = revs
    r = client.get(f"/api/revisions/{cur}/qto?base_revision_id={base}", headers=estimator)
    assert r.status_code == 200, r.text
    q = r.json()["data"]
    m = q["measures"]
    assert m["demo_wall_length_m"] == 4 and m["new_wall_length_m"] == 5
    assert m["keep_wall_length_m"] == 5 and m["new_door_count"] == 1
    assert q["base_revision_id"] == base and q["rooms"][0]["name"] == "거실"
    assert q["rules_source"] == "DEFAULT"
    # no base: everything is new work
    m = client.get(f"/api/revisions/{cur}/qto", headers=estimator).json()["data"]["measures"]
    assert m["new_wall_length_m"] == 10 and m["demo_wall_length_m"] == 0
    r = client.get(f"/api/revisions/{cur}/qto?base_revision_id={cur}", headers=estimator)
    assert r.status_code == 422 and r.json()["error"]["code"] == "QTO_BASE_INVALID"
    r = client.get(f"/api/revisions/{cur}/qto?base_revision_id=x", headers=estimator)
    assert r.status_code == 422  # not a revision id


def test_qto_base_from_another_project_is_refused(client, world, estimator, revs, tmp_path):  # noqa: F811
    _, cur = revs
    admin = world.h["admin"]
    pid = client.post(
        "/api/projects", json={"project_code": "P2", "project_name": "o"}, headers=admin
    ).json()["data"]["project_id"]
    doc = client.post(
        f"/api/projects/{pid}/documents",
        json={"doc_no": "D2", "doc_type": "DRAWING", "title": "T"},
        headers=admin,
    ).json()["data"]["document_id"]
    r = client.post(
        f"/api/documents/{doc}/dxf",
        files={"file": ("p.dxf", dxf(tmp_path, "o.dxf", [W1]))},
        headers=admin,
    )
    other = r.json()["data"]["revision_id"]
    url = f"/api/revisions/{cur}/qto?base_revision_id={other}"
    assert client.get(url, headers=estimator).status_code == 404  # not a member there
    r = client.get(url, headers=admin)
    assert r.status_code == 422 and r.json()["error"]["code"] == "QTO_BASE_INVALID"


DEMO = WORK | {
    "item_code": "D-WALL",
    "item_name": "벽체 철거",
    "work_category": "철거",
    "qty_formula": "demo_wall_length_m * height_m",
    "loss_rate": None,
    "material_price": "0",
    "labor_price": "20000",
    "expense_price": "5000",
    "valid_until": None,
}


@pytest.fixture
def active(client, world):
    admin = world.h["admin"]
    vid = client.post(
        "/api/master-versions",
        json={"version_code": "C1", "effective_from": "2020-01-01"},
        headers=admin,
    ).json()["data"]["version_id"]
    ratios = FULL["cost_ratios"] | {"indirect_labor_rate": "0.1", "indirect_cost_rate": "0.05"}
    bundle = FULL | CONSTRUCTION | {"work_items": [WORK, DEMO], "cost_ratios": ratios}
    bundle["mapping_rules"] = []  # default QTO layers (A-WALL is a wall)
    assert client.put(f"/api/master-versions/{vid}", json=bundle, headers=admin).status_code == 200
    assert client.post(f"/api/master-versions/{vid}/activate", headers=admin).status_code == 200
    return vid


def test_construction_quote(client, world, estimator, revs, active):  # noqa: F811
    base, cur = revs
    req = {
        "kind": "CONSTRUCTION",
        "revision_id": cur,
        "base_revision_id": base,
        "height_mm": "2500",
        "region": "SEOUL",
    }
    r = client.post("/api/quotes", json=req, headers=estimator)
    assert r.status_code == 200, r.text
    q = r.json()["data"]
    assert q["source_kind"] == "CONSTRUCTION" and q["revision_id"] == cur
    lines = {(ln["item_code"], ln["cost_category"]): ln for ln in q["lines"]}
    paint = lines[("A-PAINT", "MATERIAL")]
    # wall 10 m × 2.5 m × 2 sides = 50 m2, +5 % loss, 3000 × 1.05 (SEOUL)
    assert paint["calculated_qty"] == "52.500000" and paint["calculated_unit_price"] == "3150.00"
    demo = lines[("D-WALL", "LABOR")]
    assert demo["calculated_qty"] == "10.000000"  # 4 m torn down × 2.5 m
    assert demo["traces"][0]["revision_id"] == base and demo["traces"][0]["source_kind"] == "ENTITY"
    assert ("INDIRECT_LABOR", "LABOR") in lines and ("INDIRECT_COST", "OVERHEAD") in lines
    assert q["inputs"]["region"] == "SEOUL" and q["metrics"]["rooms"][0]["name"] == "거실"
    assert "items" not in q["metrics"]
    v = client.post(f"/api/quotes/{q['quote_id']}/validate", headers=estimator).json()["data"]
    errors = [lg for lg in v["logs"] if lg["severity"] == "ERROR"]
    assert errors == []  # header totals agree with the construction summary (profit basis)

    # what-if: labour +10 %, and the sensitivity ranking; nothing is stored
    r = client.post(
        f"/api/quotes/{q['quote_id']}/scenarios",
        json={"scenarios": [{"name": "노임 상승", "labor": "1.1"}]},
        headers=estimator,
    )
    assert r.status_code == 200, r.text
    s = r.json()["data"]
    assert s["base"]["totals"]["total_amount"] == q["total_amount"]
    sc = s["scenarios"][0]
    assert sc["name"] == "노임 상승" and D(sc["delta"]) > 0
    assert D(sc["totals"]["labor_cost"]) == (D(q["labor_cost"]) * D("1.1")).quantize(D("0.01"))
    assert len(s["sensitivity"]) == 10
    up = next(x for x in s["sensitivity"] if x["factor"] == "qty" and D(x["change"]) > 0)
    assert D(up["delta"]) > 0
    r = client.post(
        f"/api/quotes/{q['quote_id']}/scenarios?format=xlsx",
        json={"scenarios": [{"name": "=HYPERLINK(1)", "material": "0.9"}]},
        headers=estimator,
    )
    assert r.status_code == 200 and r.content.startswith(b"PK")
    from openpyxl import load_workbook

    assert r.headers["cache-control"] == "no-store"
    ws = load_workbook(io.BytesIO(r.content))["시나리오 비교"]
    assert ws["A4"].value == "=HYPERLINK(1)" and ws["A4"].data_type == "s"
    assert ws["A4"].quotePrefix  # stays text when edited in Excel
    # a factor that overflows a line is reported, not shown as a plain delta
    r = client.post(
        f"/api/quotes/{q['quote_id']}/scenarios",
        json={"scenarios": [{"material": "10", "labor": "10", "qty": "10"}]},
        headers=estimator,
    )
    assert r.headers["cache-control"] == "no-store"
    assert r.json()["data"]["scenarios"][0]["errors"] == []  # in range: no codes


def test_construction_quote_inputs(client, world, estimator, revs, active):  # noqa: F811
    _, cur = revs
    bad = [
        {"kind": "CONSTRUCTION", "document_id": world.doc},  # needs a drawing revision
        {"revision_id": cur, "height_mm": "2500"},  # height is for construction quotes
        {"kind": "CONSTRUCTION", "revision_id": cur, "scenario": {"labor": "0"}},
        {"kind": "CONSTRUCTION", "revision_id": cur, "region": "../x"},
    ]
    for req in bad:
        r = client.post("/api/quotes", json=req, headers=estimator)
        assert r.status_code == 422, req
    # no height: the wall-area items are errors, the quote still records them
    r = client.post(
        "/api/quotes", json={"kind": "CONSTRUCTION", "revision_id": cur}, headers=estimator
    )
    q = r.json()["data"]
    assert q["has_errors"] and any(lg["code"] == "QUOTE_HEIGHT_REQUIRED" for lg in q["logs"])
    r = client.post(
        "/api/quotes",
        json={"kind": "CONSTRUCTION", "revision_id": cur, "height_mm": "2500", "region": "BUSAN"},
        headers=estimator,
    )
    assert any(lg["code"] == "QUOTE_FACTOR_UNKNOWN" for lg in r.json()["data"]["logs"])


def test_construction_report(client, world, estimator, revs, active, supplier):  # noqa: F811
    from openpyxl import load_workbook

    base, cur = revs
    req = {
        "kind": "CONSTRUCTION",
        "revision_id": cur,
        "base_revision_id": base,
        "height_mm": "2500",
    }
    q = client.post("/api/quotes", json=req, headers=estimator).json()["data"]
    url = f"/api/quotes/{q['quote_id']}/report"
    r = client.get(url + "?format=xlsx&basis=true", headers=estimator)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    assert {"공사원가계산서", "내역서", "실별면적표"} <= set(wb.sheetnames)
    st = {row[0]: row[1] for row in wb["공사원가계산서"].iter_rows(min_row=2, values_only=True)}
    assert D(str(st["도급액 합계"])) == D(q["total_amount"])
    assert D(str(st["노무비 계"])) == D(str(st["  직접노무비"])) + D(str(st["  간접노무비"]))
    boq = list(wb["내역서"].iter_rows(min_row=2, values_only=True))
    paint = next(r for r in boq if r[1] == "벽체 도장")
    assert paint[4] == 50 and paint[5] == 0.05  # net quantity; material carries the loss
    assert sum(D(str(r[-1])) for r in boq) + D(str(st["  간접노무비"])) + D(
        str(st["  간접공사비"])
    ) == D(q["manufacturing_cost"])
    room = next(wb["실별면적표"].iter_rows(min_row=2, values_only=True))
    assert room[0] == "거실" and room[1] == 20 and room[3] == "유지"
    assert "순공사원가" in [c.value for c in wb["원가내역"]["A"]]
    basis = list(wb["산출근거"].iter_rows(min_row=3))  # labels as on the quote screen
    assert "경비" in {r[1].value for r in basis} and "제조간접비" not in {r[1].value for r in basis}
    ratio = next(r for r in basis if r[4].value == "식")
    assert ratio[5].number_format == "0.0###%" and 0 < ratio[5].value < 1  # exact rate, as %
    r = client.get(url + "?basis=true", headers=estimator)
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    text = "".join(pg.extract_text() for pg in PdfReader(io.BytesIO(r.content)).pages)
    assert "%" in text and "제조간접비" not in text
