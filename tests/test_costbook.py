"""Track 8 P4-6: 건설공사 표준품셈 in construction quotes — a work item priced from its 품셈 item
(resources per unit x resource prices = 일위대가) next to directly priced ones, traced to the
notice and page, checked at activation, shown as a 일위대가표 in the report."""

import io
import itertools
from decimal import Decimal as D

import pytest

from core.quote_engine.construction import compute_construction
from tests.test_api_construction import revs  # noqa: F401 - fixture
from tests.test_api_master import CONSTRUCTION, FULL, WORK
from tests.test_api_quote import estimator  # noqa: F401 - fixture
from tests.test_construction import QTO, RATIOS, SRC, work

NOTICE = "국토교통부 공고 제2025-1574호"
ITEM = {
    "item_code": "CB-PLASTER-W",
    "item_name": "시멘트 모르타르 미장",
    "spec": "벽, 18mm",
    "unit": "m2",
    "category": "미장",
    "notice_no": NOTICE,
    "effective_from": "2026-01-01",
    "page_no": 512,
    "source_sha256": "a" * 64,
    "extract_method": "MANUAL",
    "reviewer": "검수자",
    "reviewed_at": "2026-01-05",
}
# synthetic test numbers (not the 2026 품셈's): per m2
RESOURCES = [
    {
        "item_code": "CB-PLASTER-W",
        "resource_type": "LABOR",
        "resource_code": "L-PLASTERER",
        "qty_per_unit": "0.06",
    },
    {
        "item_code": "CB-PLASTER-W",
        "resource_type": "LABOR",
        "resource_code": "L-WORKER",
        "qty_per_unit": "0.02",
    },
    {
        "item_code": "CB-PLASTER-W",
        "resource_type": "MATERIAL",
        "resource_code": "M-CEMENT",
        "qty_per_unit": "8.5",
    },
    {
        "item_code": "CB-PLASTER-W",
        "resource_type": "EQUIPMENT",
        "resource_code": "E-MIXER",
        "qty_per_unit": "0.01",
    },
]
PRICES = [
    {
        "resource_code": "L-PLASTERER",
        "resource_type": "LABOR",
        "resource_name": "미장공",
        "unit": "인",
        "unit_price": "280000",
        "price_date": "2026-01-01",
        "source": "시중노임단가",
    },
    {
        "resource_code": "L-WORKER",
        "resource_type": "LABOR",
        "resource_name": "보통인부",
        "unit": "인",
        "unit_price": "170000",
    },
    {
        "resource_code": "M-CEMENT",
        "resource_type": "MATERIAL",
        "resource_name": "시멘트",
        "unit": "kg",
        "unit_price": "150",
    },
    {
        "resource_code": "E-MIXER",
        "resource_type": "EQUIPMENT",
        "resource_name": "믹서",
        "unit": "시간",
        "unit_price": "20000",
    },
]
PLASTER = {
    **work("A-PLASTER", "wall_length_m * height_m * 2", None, None, None, "0.05"),
    "item_name": "벽 미장",
    "costbook_item_code": "CB-PLASTER-W",
}


def test_unit_cost_lines_from_the_costbook():
    bundle = {
        "cost_ratios": RATIOS,
        "work_items": [PLASTER],
        "costbook_items": [ITEM],
        "costbook_resources": RESOURCES,
        "resource_prices": PRICES,
    }
    r = compute_construction(QTO, bundle, SRC, height_m=D("2.5"))
    assert not r["has_errors"]
    by = {ln["item_code"]: ln for ln in r["lines"]}
    # 50 m2 of wall faces (10 m x 2.5 m x 2)
    assert by["A-PLASTER/L-PLASTERER"]["calculated_amount"] == D("840000.00")  # 3 인 x 280,000
    assert by["A-PLASTER/L-WORKER"]["calculated_amount"] == D("170000.00")
    cement = by["A-PLASTER/M-CEMENT"]  # 50 x 8.5 kg x 1.05 loss
    assert cement["calculated_qty"] == D("446.25") and cement["calculated_amount"] == D("66937.50")
    mixer = by["A-PLASTER/E-MIXER"]  # equipment = 경비
    assert mixer["cost_category"] == "OVERHEAD" and mixer["calculated_amount"] == D("10000.00")
    assert by["A-PLASTER/L-PLASTERER"]["item_name"] == "벽 미장 · 미장공"
    t = by["A-PLASTER/L-PLASTERER"]["traces"][0]
    assert t["rule_code"] == "COSTBOOK:CB-PLASTER-W" and t["sources"] == ["A1", "A2"]
    assert t["inputs"]["notice_no"] == NOTICE and t["inputs"]["page_no"] == 512
    assert t["inputs"]["price_source"] == "시중노임단가" and D(t["inputs"]["work_qty"]) == 50
    assert r["totals"]["labor_cost"] == D("1111000.00")  # 1,010,000 + 10 % indirect


def test_missing_costbook_rows_are_errors():
    bundle = {"cost_ratios": RATIOS, "work_items": [PLASTER], "costbook_items": [ITEM]}
    r = compute_construction(QTO, bundle, SRC, height_m=D("2.5"))
    assert any(lg["code"] == "QUOTE_COSTBOOK_EMPTY" for lg in r["logs"]) and r["has_errors"]


CODES = itertools.count(1)


def master(client, admin, bundle):
    vid = client.post(
        "/api/master-versions",
        json={"version_code": f"CB{next(CODES)}", "effective_from": "2020-01-01"},
        headers=admin,
    ).json()["data"]["version_id"]
    return vid, client.put(f"/api/master-versions/{vid}", json=bundle, headers=admin)


def test_costbook_master_data_rules(client, world):
    admin = world.h["admin"]
    ratios = FULL["cost_ratios"] | {"indirect_labor_rate": "0.1", "indirect_cost_rate": "0.05"}
    plaster = {
        **WORK,
        "item_code": "A-PLASTER",
        "material_price": None,
        "labor_price": None,
        "expense_price": None,
        "costbook_item_code": "CB-PLASTER-W",
    }
    base = FULL | CONSTRUCTION | {"cost_ratios": ratios, "work_items": [plaster]}
    full = base | {
        "costbook_items": [ITEM],
        "costbook_resources": RESOURCES,
        "resource_prices": PRICES,
    }
    # a work item with both its own price and a 품셈 item, or a resource of an unknown item
    _, r = master(client, admin, full | {"work_items": [plaster | {"labor_price": "1"}]})
    assert r.status_code == 422
    _, r = master(client, admin, full | {"costbook_resources": [RESOURCES[0] | {"item_code": "X"}]})
    assert r.status_code == 422
    # activation needs a price for every resource the linked item uses
    vid, r = master(client, admin, full | {"resource_prices": PRICES[:3]})
    assert r.status_code == 200, r.text
    r = client.post(f"/api/master-versions/{vid}/activate", headers=admin)
    assert (
        r.status_code == 422
        and "resource_prices.E-MIXER.unit_price" in r.json()["error"]["message"]
    )


@pytest.fixture
def costbook(client, world):
    admin = world.h["admin"]
    ratios = FULL["cost_ratios"] | {"indirect_labor_rate": "0.1", "indirect_cost_rate": "0.05"}
    plaster = {
        **WORK,
        "item_code": "A-PLASTER",
        "item_name": "벽 미장",
        "material_price": None,
        "labor_price": None,
        "expense_price": None,
        "loss_rate": "0.05",
        "costbook_item_code": "CB-PLASTER-W",
    }
    bundle = (
        FULL
        | CONSTRUCTION
        | {
            "cost_ratios": ratios,
            "work_items": [WORK, plaster],
            "mapping_rules": [],
            "costbook_items": [ITEM],
            "costbook_resources": RESOURCES,
            "resource_prices": PRICES,
        }
    )
    vid, r = master(client, admin, bundle)
    assert r.status_code == 200, r.text
    assert client.post(f"/api/master-versions/{vid}/activate", headers=admin).status_code == 200


def test_construction_quote_with_unit_costs(client, world, estimator, revs, costbook):  # noqa: F811
    from openpyxl import load_workbook

    base, cur = revs
    req = {
        "kind": "CONSTRUCTION",
        "revision_id": cur,
        "base_revision_id": base,
        "height_mm": "2500",
    }
    q = client.post("/api/quotes", json=req, headers=estimator).json()["data"]
    codes = {ln["item_code"] for ln in q["lines"]}
    assert {"A-PAINT", "A-PLASTER/L-PLASTERER", "A-PLASTER/M-CEMENT"} <= codes  # side by side
    v = client.post(f"/api/quotes/{q['quote_id']}/validate", headers=estimator).json()["data"]
    assert not any(lg["code"] == "QUOTE_DUPLICATE_LINE" for lg in v["logs"])
    r = client.get(f"/api/quotes/{q['quote_id']}/report?format=xlsx", headers=estimator)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    boq = {row[1]: row for row in wb["내역서"].iter_rows(min_row=2, values_only=True)}
    plaster = boq["벽 미장"]  # one row per work item: 50 m2, labour 1,010,000 / 50
    assert plaster[4] == 50 and D(str(plaster[8])) == D(20200)
    uc = list(wb["일위대가표"].iter_rows(min_row=2, values_only=True))
    worker = next(r for r in uc if r[5] == "보통인부")
    assert worker[2] == "CB-PLASTER-W" and worker[3] == NOTICE and worker[4] == 512
    assert D(str(worker[8])) == D("0.02") and D(str(worker[10])) == D(3400)  # per m2
    pdf = client.get(f"/api/quotes/{q['quote_id']}/report", headers=estimator)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


def test_drawing_object_measures_trace_to_their_entities():
    qto = {
        "measures": {"obj_slab": 4.04, "new_obj_slab": 4.04},
        "items": [{"role": "OBJ:slab", "status": "NEW", "handles": ["S1"]}],
    }
    bundle = {"cost_ratios": RATIOS, "work_items": [work("C-SLAB", "obj_slab", "100000", "0", "0")]}
    r = compute_construction(qto, bundle, SRC, height_m=None)
    line = r["lines"][0]
    assert line["calculated_qty"] == D("4.04") and line["traces"][0]["sources"] == ["S1"]


def test_drawing_objects_in_master_data_and_quotes(client, world, estimator, revs):  # noqa: F811
    admin = world.h["admin"]
    obj = {
        "code": "wall2",
        "name": "벽(객체)",
        "layer_pattern": "A-WALL",
        "quantity_type": "LENGTH",
        "factor": "1",
        "waste_rate": "0.1",
    }
    item = {**WORK, "item_code": "A-OBJ", "item_name": "객체 벽", "qty_formula": "new_obj_wall2"}
    ratios = FULL["cost_ratios"] | {"indirect_labor_rate": "0.1", "indirect_cost_rate": "0.05"}
    base = FULL | CONSTRUCTION | {"cost_ratios": ratios, "mapping_rules": [], "work_items": [item]}
    _, r = master(client, admin, base)  # obj_wall2 without the object
    assert r.status_code == 422
    _, r = master(client, admin, base | {"qto_objects": [obj | {"layer_pattern": None}]})
    assert r.status_code == 422  # neither a layer nor a block
    vid, r = master(client, admin, base | {"qto_objects": [obj]})
    assert r.status_code == 200, r.text
    assert r.json()["data"]["qto_objects"][0]["code"] == "wall2"
    r = client.post(f"/api/master-versions/{vid}/activate", headers=admin)
    assert r.status_code == 200, r.text
    b, cur = revs
    m = client.get(f"/api/revisions/{cur}/qto?base_revision_id={b}", headers=estimator).json()
    assert m["data"]["measures"]["new_obj_wall2"] == pytest.approx(5.5)  # W3 5 m x 1.1
    req = {"kind": "CONSTRUCTION", "revision_id": cur, "base_revision_id": b}
    q = client.post("/api/quotes", json=req, headers=estimator).json()["data"]
    ln = next(
        ln for ln in q["lines"] if ln["item_code"] == "A-OBJ" and ln["cost_category"] == "LABOR"
    )
    assert D(ln["calculated_qty"]) == D("5.5") and len(ln["traces"][0]["sources"]) == 1


EXTERNAL = {  # a Revit / Dynamo export (doc example shape; extra keys are ignored)
    "project_id": "P-2026-001",
    "source_file": "sample.rvt",
    "elements": [
        {
            "element_id": "331245",
            "category": "Walls",
            "family": "Basic Wall",
            "quantity_type": "AREA",
            "quantity": 30,
            "unit": "m2",
            "level": "1F",
            "work_item_code": "A-PLASTER",
        },
        {"element_id": "331246", "quantity": 20, "unit": "m2", "work_item_code": "A-PLASTER"},
        {"element_id": "9", "quantity": 1, "unit": "m3", "work_item_code": "X-NONE"},
    ],
}


def test_external_quantities_replace_the_formula():
    bundle = {
        "cost_ratios": RATIOS,
        "work_items": [PLASTER],
        "costbook_items": [ITEM],
        "costbook_resources": RESOURCES,
        "resource_prices": PRICES,
    }
    qto = {**QTO, "external": EXTERNAL}
    r = compute_construction(qto, bundle, SRC, height_m=None)  # no height: the formula is unused
    assert not r["has_errors"]
    by = {ln["item_code"]: ln for ln in r["lines"]}
    assert by["A-PLASTER/L-PLASTERER"]["calculated_amount"] == D("840000.00")  # 50 m2 again
    t = by["A-PLASTER/L-PLASTERER"]["traces"][0]
    assert t["source_kind"] == "EXTERNAL" and t["sources"] == ["331245", "331246"]
    assert t["inputs"]["source_file"] == "sample.rvt" and D(t["inputs"]["work_qty"]) == 50
    assert any(lg["code"] == "QUOTE_EXTERNAL_UNMAPPED" for lg in r["logs"])
    bad = {**EXTERNAL, "elements": [EXTERNAL["elements"][0] | {"unit": "m"}]}
    r = compute_construction({**QTO, "external": bad}, bundle, SRC, height_m=None)
    assert any(lg["code"] == "QUOTE_UNIT_MISMATCH" for lg in r["logs"]) and r["has_errors"]


def test_construction_quote_with_external_quantities(client, world, estimator, revs, costbook):  # noqa: F811
    _, cur = revs
    req = {"kind": "CONSTRUCTION", "revision_id": cur, "height_mm": "2500", "quantities": EXTERNAL}
    r = client.post("/api/quotes", json=req, headers=estimator)
    assert r.status_code == 200, r.text
    q = r.json()["data"]
    ln = next(ln for ln in q["lines"] if ln["item_code"] == "A-PLASTER/L-WORKER")
    assert D(ln["calculated_qty"]) == 1  # 50 m2 x 0.02 인
    assert ln["traces"][0]["source_kind"] == "EXTERNAL"
    assert q["inputs"]["quantities_file"] == "sample.rvt"
    # a scenario recomputes from the stored quantities
    s = client.post(
        f"/api/quotes/{q['quote_id']}/scenarios",
        json={"scenarios": [{"labor": "1.1"}]},
        headers=estimator,
    )
    assert s.status_code == 200, s.text
    hostile = {**EXTERNAL, "elements": [EXTERNAL["elements"][0] | {"element_id": "a\x00b"}]}
    r = client.post("/api/quotes", json=req | {"quantities": hostile}, headers=estimator)
    assert r.status_code == 422
    r = client.post(
        "/api/quotes",
        json={"kind": "PART", "revision_id": cur, "quantities": EXTERNAL},
        headers=estimator,
    )
    assert r.status_code == 422
