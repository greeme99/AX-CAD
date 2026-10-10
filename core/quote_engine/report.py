"""FN-21 quote document (Korean standard layout) as PDF (reportlab, embedded NanumGothic) or XLSX.

Input is a plain dict built by the API (`backend/api/routes_quote.py::_report_model`); nothing here
touches the database. User-supplied text (titles, reasons, customer names) is escaped for the PDF
paragraph markup and forced to plain strings in the workbook (no formula injection).
"""

import io
import re
import threading
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

FONT_DIR = Path(__file__).resolve().parents[2] / "backend" / "assets" / "fonts"
TOTAL_LABELS = {
    "material_cost": "직접재료비",
    "labor_cost": "직접노무비",
    "overhead_cost": "제조간접비",
    "outsource_cost": "외주가공비",
    "manufacturing_cost": "순제조원가",
    "admin_cost": "일반관리비",
    "total_cost": "총원가",
    "profit": "이윤",
    "supply_amount": "공급가액",
    "vat_amount": "세액(VAT)",
    "total_amount": "합계",
}
CATEGORY = {
    "MATERIAL": "재료비",
    "LABOR": "노무비",
    "OVERHEAD": "제조간접비",
    "OUTSOURCE": "외주가공비",
}
DIGITS = "영일이삼사오육칠팔구"
CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")  # openpyxl refuses them, PDFs garble
FORMULA_LEAD = ("=", "+", "-", "@")
INTERNAL = "내부용·대외비"
_font_lock = threading.Lock()


def clean(v: Any) -> str:
    return CONTROL.sub("", "" if v is None else str(v))


def korean_amount(n: int) -> str:
    """62930 -> '육만이천구백삼십' (formal: 일 is always written, 10000 -> '일만')."""
    if n < 0:
        raise ValueError("negative amount")
    if n == 0:
        return "영"
    groups = []
    for big in ("", "만", "억", "조", "경"):
        n, chunk = divmod(n, 10000)
        if chunk:
            s = "".join(
                DIGITS[d] + unit
                for d, unit in zip(
                    (chunk // 1000, chunk // 100 % 10, chunk // 10 % 10, chunk % 10),
                    ("천", "백", "십", ""),
                    strict=True,
                )
                if d
            )
            groups.append(s + big)
        if not n:
            break
    if n:
        raise ValueError("amount too large")
    return "".join(reversed(groups))


def won(v: Any, digits: int = 0) -> str:
    if v is None:
        return "-"
    q = Decimal(1).scaleb(-digits)
    d = Decimal(str(v)).quantize(q, rounding=ROUND_HALF_UP)
    return f"{d:,.{digits}f}"


def price(v: Any) -> str:
    """Unit price: whole won as is, a fraction shown to the jeon so qty x price = supply."""
    return won(v, 0 if Decimal(str(v)) % 1 == 0 else 2)


def _qty(v: Any) -> str:
    d = Decimal(str(v)).normalize()
    return f"{d:,f}"


# --- PDF ------------------------------------------------------------------------------------


def _fonts() -> None:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    with _font_lock:  # first concurrent renders must not register twice
        if "Nanum-Bold" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("Nanum", FONT_DIR / "NanumGothic-Regular.ttf"))
            pdfmetrics.registerFont(TTFont("Nanum-Bold", FONT_DIR / "NanumGothic-Bold.ttf"))


def render_pdf(m: dict[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table
    from reportlab.platypus import TableStyle as TS

    _fonts()
    base = ParagraphStyle("b", fontName="Nanum", fontSize=9, leading=12)
    small = ParagraphStyle("s", parent=base, fontSize=7.5, leading=9.5)
    right = ParagraphStyle("r", parent=base, alignment=TA_RIGHT)
    bold = ParagraphStyle("bd", parent=base, fontName="Nanum-Bold")
    title = ParagraphStyle("t", parent=bold, fontSize=22, leading=28, alignment=TA_CENTER)
    h2 = ParagraphStyle("h2", parent=bold, fontSize=11, leading=15, spaceBefore=6, spaceAfter=4)

    def p(text: Any, style: ParagraphStyle = base) -> Paragraph:
        return Paragraph(escape(clean(text)), style)

    grid = [
        ("FONTNAME", (0, 0), (-1, -1), "Nanum"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#9aa4b2")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    head = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef1f5"))]

    s, c, it = m["supplier"], m["customer"], m["item"]
    total = Decimal(str(m["totals"]["total_amount"])).quantize(1, rounding=ROUND_HALF_UP)
    story: list[Any] = [p("견 적 서", title), Spacer(1, 4 * mm)]

    left = [
        [p("견적번호", bold), p(m["quote_no"])],
        [p("견적일", bold), p(m["issued"])],
        [p("유효기간", bold), p(f"{m['valid_until']}까지")],
        [p("수신", bold), p(f"{c['name']} 귀하")],
        [p("건명", bold), p(c["project"])],
    ]
    sup = [
        [p("공급자", bold), p("등록번호", bold), p(s["business_no"])],
        ["", p("상호", bold), p(s["company"])],
        ["", p("대표자", bold), p(s["ceo"])],
        ["", p("주소", bold), p(s["address"])],
        ["", p("연락처", bold), p(f"{s['phone']} / {s['email']}")],
    ]
    t_left = Table(left, colWidths=[18 * mm, 62 * mm])
    t_left.setStyle(TS(grid))
    t_sup = Table(sup, colWidths=[12 * mm, 18 * mm, 60 * mm])
    t_sup.setStyle(TS([*grid, ("SPAN", (0, 0), (0, -1))]))
    story += [Table([[t_left, t_sup]], colWidths=[84 * mm, 92 * mm]), Spacer(1, 4 * mm)]
    story.append(p("아래와 같이 견적합니다."))
    story.append(p(f"합계금액(VAT 포함): 일금 {korean_amount(int(total))}원정 (₩{won(total)})", h2))

    rows = [
        [p(h, bold) for h in ("No", "품명", "규격", "수량", "단위", "단가", "공급가액", "세액")],
        [
            p("1"),
            p(it["name"]),
            p(it["spec"]),
            p(_qty(it["qty"]), right),
            p(it["unit"]),
            p(price(it["unit_price"]), right),
            p(won(it["supply"]), right),
            p(won(it["vat"]), right),
        ],
        [
            "",
            p("합계", bold),
            "",
            "",
            "",
            "",
            p(won(it["supply"]), right),
            p(won(it["vat"]), right),
        ],
    ]
    widths = [10, 42, 36, 14, 12, 20, 24, 18]
    t = Table(rows, colWidths=[w * mm for w in widths])
    t.setStyle(TS([*grid, *head, ("SPAN", (1, 2), (5, 2))]))
    story += [t, Spacer(1, 4 * mm)]

    terms = [
        [p("납기", bold), p(m["terms"]["delivery"])],
        [p("결제조건", bold), p(m["terms"]["payment"])],
        [p("비고", bold), p(m["terms"]["note"])],
    ]
    tt = Table(terms, colWidths=[24 * mm, 152 * mm])
    tt.setStyle(TS(grid))
    story += [tt, Spacer(1, 3 * mm)]
    story.append(
        p(
            f"AX-CAD 자동 산출 · 기준정보 #{m['master_version_id']} · "
            f"원천 {m['source']} · 수동 조정 {m['manual_count']}건",
            small,
        )
    )

    if m.get("basis"):
        story += [PageBreak(), p(f"[{INTERNAL}] 별첨 1. 원가 내역", h2)]
        cost = [[p("항목", bold), p("금액(원)", bold)]] + [
            [p(label), p(won(m["totals"][k]), right)] for k, label in TOTAL_LABELS.items()
        ]
        ct = Table(cost, colWidths=[60 * mm, 50 * mm])
        ct.setStyle(TS([*grid, *head]))
        story += [ct, p(f"[{INTERNAL}] 별첨 2. 산출근거", h2)]
        basis = [[p(h, bold) for h in ("No", "구분", "항목", "수량", "단가", "금액", "근거")]] + [
            [
                p(ln["line_no"], small),
                p(CATEGORY.get(ln["cost_category"], ln["cost_category"]), small),
                p(ln["item_name"], small),
                p(f"{_qty(ln['qty'])} {ln['unit']}", small),
                p(won(ln["unit_price"], 2), small),
                p(won(ln["amount"]), small),
                p(
                    (f"[수동] {ln['override_reason']} / " if ln["manual"] else "") + ln["basis"],
                    small,
                ),
            ]
            for ln in m["lines"]
        ]
        bt = Table(basis, colWidths=[w * mm for w in (10, 18, 28, 22, 20, 20, 58)], repeatRows=1)
        bt.setStyle(TS([*grid, *head]))
        story.append(bt)

    def deco(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont("Nanum", 7)
        canvas.drawRightString(A4[0] - 15 * mm, 10 * mm, f"{m['quote_no']} · {doc.page}")
        if not m["official"]:  # FN-21: anything before approval is visibly a draft
            canvas.setFont("Nanum-Bold", 64)
            canvas.setFillColor(colors.Color(0.8, 0.1, 0.1, alpha=0.18))
            canvas.translate(A4[0] / 2, A4[1] / 2)
            canvas.rotate(35)
            canvas.drawCentredString(0, 0, "초안(DRAFT)")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=17 * mm,
        rightMargin=17 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=f"견적서 {m['quote_no']}",
        author=s["company"],
        invariant=1,  # no creation timestamp/ID: same data -> same bytes
    )
    doc.build(story, onFirstPage=deco, onLaterPages=deco)
    return buf.getvalue()


# --- XLSX -----------------------------------------------------------------------------------


def render_xlsx(m: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "견적서"

    def put(row: list[Any], sheet: Any = ws) -> None:
        sheet.append([clean(v) if isinstance(v, str) else v for v in row])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # "=..." from a title or reason stays text, never a formula
                # and stays text when someone edits the cell or copies it out
                cell.quotePrefix = cell.value.startswith(FORMULA_LEAD)

    s, c, it = m["supplier"], m["customer"], m["item"]
    put(["견 적 서" if m["official"] else "견 적 서 (초안 DRAFT — 승인 전, 외부 발송 불가)"])
    ws["A1"].font = Font(bold=True, size=16)
    for row in (
        ["견적번호", m["quote_no"], "", "공급자 등록번호", s["business_no"]],
        ["견적일", m["issued"], "", "상호", s["company"]],
        ["유효기간", f"{m['valid_until']}까지", "", "대표자", s["ceo"]],
        ["수신", f"{c['name']} 귀하", "", "주소", s["address"]],
        ["건명", c["project"], "", "연락처", f"{s['phone']} / {s['email']}"],
    ):
        put(row)
    total = Decimal(str(m["totals"]["total_amount"])).quantize(1, rounding=ROUND_HALF_UP)
    put([])
    put(["합계금액(VAT 포함)", f"일금 {korean_amount(int(total))}원정", total])
    put([])
    put(["No", "품명", "규격", "수량", "단위", "단가", "공급가액", "세액"])
    put(
        [1, it["name"], it["spec"], Decimal(str(it["qty"])), it["unit"]]
        + [Decimal(str(it[k])) for k in ("unit_price", "supply", "vat")]
    )
    put(["", "합계", "", "", "", "", Decimal(str(it["supply"])), Decimal(str(it["vat"]))])
    put([])
    for k, label in (("delivery", "납기"), ("payment", "결제조건"), ("note", "비고")):
        put([label, m["terms"][k]])
    for col, w in zip("ABCDEFGH", (14, 30, 22, 16, 30, 14, 14, 12), strict=True):
        ws.column_dimensions[col].width = w

    if m.get("basis"):
        cs = wb.create_sheet("원가내역")
        cs.append([INTERNAL])
        cs.append(["항목", "금액(원)"])
        for k, label in TOTAL_LABELS.items():
            cs.append([label, Decimal(str(m["totals"][k]))])
        bs = wb.create_sheet("산출근거")
        bs.append([INTERNAL])
        bs.append(
            ["No", "구분", "항목", "수량", "단위", "단가", "금액", "수동", "조정 사유", "근거"]
        )
        for ln in m["lines"]:
            put(
                [
                    ln["line_no"],
                    CATEGORY.get(ln["cost_category"], ln["cost_category"]),
                    ln["item_name"],
                    Decimal(str(ln["qty"])),
                    ln["unit"],
                    None if ln["unit_price"] is None else Decimal(str(ln["unit_price"])),
                    None if ln["amount"] is None else Decimal(str(ln["amount"])),
                    "Y" if ln["manual"] else "",
                    ln["override_reason"] or "",
                    ln["basis"],
                ],
                bs,
            )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
