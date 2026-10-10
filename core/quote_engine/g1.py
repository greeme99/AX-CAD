"""G1 input workbook: the template generator and the importer share the sheet/column definitions
below, so docs/inputs/G1_입력양식.xlsx and POST /api/master-versions/{id}/import-xlsx never drift.

Sheets 1-6 become one master-data bundle (FN-16). Sheets 7-8 are test answer keys (TC-60/67) and
are not imported. Example rows carry "예시" in 비고 and are skipped."""

import re
import zipfile
from decimal import Decimal, InvalidOperation
from typing import Any

HEADER_ROW, FIRST_ROW = 4, 5
EXAMPLE = "예시"
MAX_ROWS = 2000  # per sheet, same as the API list limits
MAX_UNZIPPED = 50 * 1024**2  # xlsx is a zip: bound what openpyxl may inflate
MAX_ZIP_ENTRIES = 500

# sheet -> (description, [(header, width, note)], example rows)
S_MAT, S_PROC, S_LAB, S_RATE, S_LAYER, S_TITLE, S_SAMPLE, S_QUOTE = (
    "1_재질",
    "2_공정",
    "3_임률",
    "4_원가비율",
    "5_레이어",
    "6_표제란",
    "7_샘플기대값",
    "8_수기견적",
)
METRICS = (
    "cutting_length_mm, hole_count, punch_hole_count, bend_count, bend_length_mm, net_area_mm2, "
    "surface_area_mm2, volume_mm3, part_qty"
)
SHEETS: dict[str, tuple[str, list[tuple[str, int, str | None]], list[tuple[Any, ...]]]] = {
    S_MAT: (
        "재질 마스터(material_master). 단가는 기준정보 버전으로 관리됩니다. 비중은 재질 규격값 그대로 두셔도 됩니다.",
        [
            ("재질코드", 12, "도면 표제란에 적히는 표기 그대로(예: SS400, SUS304). 영문·숫자·.-_"),
            ("재질명", 18, None),
            ("비중(g/cm³)", 12, "중량 = 순면적 × 두께 × 비중 (판재) 또는 부피 × 비중 (3D)"),
            ("단가(원/kg)", 12, "사내 구매 단가"),
            ("스크랩률(%)", 12, "직접재료비 = 정미중량 × 단가 × (1 + 스크랩률)"),
            (
                "두께 범위(mm)",
                14,
                "두께별 단가가 다르면 행을 나눠 '1.0~3.2'처럼 기재. 전부면 '전체'",
            ),
            ("비고", 30, "'예시'로 시작하는 행은 가져오지 않습니다"),
        ],
        [
            ("SS400", "일반구조용 압연강", 7.85, None, None, "전체", EXAMPLE),
            ("SUS304", "스테인리스", 7.93, None, None, "전체", EXAMPLE),
            ("AL6061", "알루미늄 합금", 2.70, None, None, "전체", EXAMPLE),
        ],
    ),
    S_PROC: (
        "공정 규칙(process_rule). 표준공수 공식은 아래 메트릭 이름과 파라미터로 적어 주세요. 정확한 식이 없으면 비고에 현재 계산 방법을 문장으로 적어도 됩니다.",
        [
            ("공정코드", 12, "영문·숫자·.-_ (예: LASER)"),
            ("공정명", 14, None),
            ("입력 메트릭", 22, f"다음 중 하나: {METRICS}"),
            ("표준공수 공식(분)", 36, "예: setup_min + cutting_length_mm / speed_mm_per_min"),
            ("파라미터 값", 30, "예: setup_min=5, speed_mm_per_min=3000"),
            (
                "기계경비(원/h)",
                12,
                "제조간접비 = Σ 기계시간 × 기계경비 (배부 기준이 '기계시간'일 때)",
            ),
            ("임률 구분", 12, "3_임률 시트의 구분명"),
            ("비고", 30, "'예시'로 시작하는 행은 가져오지 않습니다"),
        ],
        [
            (
                "LASER",
                "레이저 절단",
                "cutting_length_mm",
                "setup_min + cutting_length_mm / speed_mm_per_min",
                "setup_min=?, speed_mm_per_min=?",
                None,
                "레이저",
                EXAMPLE,
            ),
            (
                "PUNCH",
                "NCT 펀칭",
                "punch_hole_count",
                "setup_min + punch_hole_count * sec_per_hit / 60",
                "setup_min=?, sec_per_hit=?",
                None,
                "NCT",
                EXAMPLE,
            ),
            (
                "BEND",
                "절곡",
                "bend_count",
                "setup_min + bend_count * min_per_bend",
                "setup_min=?, min_per_bend=?",
                None,
                "절곡",
                EXAMPLE,
            ),
            (
                "PAINT",
                "분체도장",
                "surface_area_mm2",
                "surface_area_mm2 / 1e6 * won_per_m2",
                "won_per_m2=?",
                None,
                None,
                EXAMPLE + ": 면적 단가형",
            ),
        ],
    ),
    S_LAB: (
        "직접노무비 = 표준공수(h) × 임률. 공정 또는 직무별로 적어 주세요.",
        [
            ("구분", 14, "2_공정 시트의 '임률 구분'과 같은 이름"),
            ("임률(원/h)", 12, None),
            ("비고", 40, "'예시'로 시작하는 행은 가져오지 않습니다"),
        ],
        [("레이저", None, EXAMPLE), ("NCT", None, EXAMPLE), ("절곡", None, EXAMPLE)],
    ),
    S_RATE: (
        "총원가 = (직접재료비 + 직접노무비 + 제조간접비) × (1 + 일반관리비율) → 공급가액 = 총원가 × (1 + 이윤율) → 합계 = 공급가액 × (1 + VAT). '값' 칸만 채워 주세요.",
        [
            ("항목", 22, None),
            ("값", 16, None),
            ("선택지 / 허용 범위", 40, None),
            ("비고", 30, None),
        ],
        [
            ("제조간접비 배부 기준", None, "기계시간 / 노무비 비율", ""),
            ("제조간접비율(%)", None, "배부 기준이 '노무비 비율'일 때만", ""),
            ("일반관리비율(%)", None, "통상 5~8", ""),
            ("이윤율(%)", None, "통상 7~15", ""),
            ("VAT(%)", 10, "고정", ""),
            ("절사 규칙", None, "절사 / 반올림 / 올림", "금액 끝자리 처리"),
            ("절사 단위(원)", None, "1 / 10 / 100 / 1000", ""),
            ("절사 적용 시점", None, "라인별 / 합계", "결과 금액이 달라지므로 꼭 선택"),
        ],
    ),
    S_LAYER: (
        "도면 레이어(또는 선종)로 엔티티 용도를 구분합니다. 와일드카드 *, 여러 개는 쉼표로 구분(예: BEND*, FOLD). 펀칭 기준 직경 행의 '레이어명 패턴' 칸에 직경(mm)을 적어 주세요.",
        [
            ("용도", 18, "외곽선(절단) / 절곡선 / 무시 / 표제란 / 펀칭 기준 직경(mm)"),
            ("레이어명 패턴", 24, None),
            ("선종(linetype)", 16, "레이어 대신 선종으로 구분하면 기재(예: DASHED = 절곡선)"),
            ("비고", 40, "'예시'로 시작하는 행은 가져오지 않습니다"),
        ],
        [
            ("외곽선(절단)", "0, OUTLINE, CUT*", None, EXAMPLE + ": 절단길이·면적 계산 대상"),
            ("절곡선", "BEND*", "DASHED", EXAMPLE),
            ("무시", "DIM*, TEXT*, CENTER*, HIDDEN*", "CENTER*, HIDDEN*", EXAMPLE),
            ("표제란", "TITLE*, FORMAT*", None, EXAMPLE),
            ("펀칭 기준 직경(mm)", 6, None, EXAMPLE + ": 이 직경 이하 원은 펀칭, 초과는 레이저"),
        ],
    ),
    S_TITLE: (
        "표제란에서 값을 읽는 규칙. 블록 속성(ATTRIB) 태그가 가장 정확합니다. 없으면 표제란에 적힌 라벨 문구를 적어 주세요(쉼표로 여러 개).",
        [
            ("항목", 12, "품번 / 품명 / 재질 / 두께 / 수량"),
            ("블록 속성 태그", 18, "표제란 블록의 ATTDEF 태그명"),
            ("문자 키워드", 30, "TEXT/MTEXT 라벨(쉼표로 여러 개)"),
            ("비고", 30, "'예시'로 시작하는 행은 가져오지 않습니다"),
        ],
        [
            ("품번", "PART_NO", "품번, PART NO, DWG NO", EXAMPLE),
            ("품명", "PART_NAME", "품명, NAME, TITLE", EXAMPLE),
            ("재질", "MATERIAL", "재질, MAT'L, MATERIAL", EXAMPLE),
            ("두께", "THICKNESS", "두께, T, THK", EXAMPLE + ": 단위 mm"),
            ("수량", "QTY", "수량, Q'TY, QTY", EXAMPLE),
        ],
    ),
    S_SAMPLE: (
        "샘플 도면별로 사람이 직접 잰 기대값. S8 회귀 테스트(TC-60, 허용오차 0.01mm)의 정답지입니다. 도면 파일은 익명화해서 함께 전달해 주세요. (기준정보로 가져오지 않음)",
        [
            ("파일명", 22, None),
            ("형식", 8, "DXF / STEP"),
            ("품번", 14, None),
            ("재질", 10, None),
            ("두께(mm)", 9, None),
            ("수량", 7, None),
            ("절단길이(mm)", 13, "외곽+내곽 윤곽 길이 합"),
            ("구멍 수", 8, None),
            ("구멍 직경별", 18, "예: Ø6×4, Ø10×2"),
            ("절곡 수", 8, None),
            ("절곡 길이(mm)", 12, None),
            ("순면적(mm²)", 13, "외곽 면적 − 구멍 면적"),
            ("STEP 파트 수", 11, "STEP만"),
            ("STEP 인스턴스 수", 12, "STEP만"),
            ("부피(mm³)", 13, "STEP만"),
            ("측정자", 10, None),
            ("비고", 24, None),
        ],
        [
            (
                "BRK-001.dxf",
                "DXF",
                "AX-1001-01",
                "SS400",
                3.2,
                4,
                1240.0,
                6,
                "Ø6×4, Ø10×2",
                2,
                180.0,
                52000.0,
                None,
                None,
                None,
                "홍길동",
                EXAMPLE,
            )
        ],
    ),
    S_QUOTE: (
        "7_샘플기대값 중 3건 이상에 대해 현재 방식으로 낸 견적 금액. 원 단위까지 일치해야 수락됩니다(TC-67). 견적서 원본도 함께 주세요. (기준정보로 가져오지 않음)",
        [
            (h, 12, None)
            for h in (
                "파일명",
                "수량",
                "직접재료비",
                "직접노무비",
                "제조간접비",
                "일반관리비",
                "이윤",
                "공급가액",
                "VAT",
                "합계",
                "견적 담당",
                "비고",
            )
        ],
        [],
    ),
}
GUIDE = [
    ("AX-CAD G1 입력 양식 (S8 메트릭·기준정보 / S9 원가 계산용)", True),
    ("", False),
    ("작성 방법", True),
    (
        "• 회색 행은 예시입니다. 비고가 '예시'로 시작하는 행은 가져오지 않으니, 덮어쓸 때는 비고의 '예시'를 지워 주세요.",
        False,
    ),
    (
        "• 노란 행 = 입력 칸. 헤더에 마우스를 올리면 설명(메모)이 보입니다. 모르는 값은 비워 두고 비고에 질문을 적어 주세요.",
        False,
    ),
    ("• 단위: 길이 mm, 면적 mm², 부피 mm³, 중량 kg, 금액 원, 비율 %(0~100).", False),
    (
        "• 시트 이름·헤더(4행)는 바꾸지 마세요. 기준정보 화면의 'G1 엑셀 가져오기'가 이 이름으로 읽습니다.",
        False,
    ),
    (
        "• 단가·임률·비율은 사내 기밀입니다. 작성본은 저장소에 커밋하지 말고 기준정보 화면에서 직접 가져오세요.",
        False,
    ),
    ("", False),
    ("제출 체크리스트 (G1 통과 조건, plan.md §2.2)", True),
    ("☐ 1_재질: 사용 재질 전부 + 원/kg 단가", False),
    ("☐ 2_공정: 레이저·펀칭·절곡·표면처리 등 공정별 표준공수 기준", False),
    ("☐ 3_임률: 공정/직무별 원/h", False),
    ("☐ 4_원가비율: 제조간접비 배부 기준, 일반관리비율, 이윤율, 절사 규칙", False),
    ("☐ 5_레이어: 도면 레이어 명명 규칙(외곽선/절곡선/무시)", False),
    ("☐ 6_표제란: 품번·품명·재질·두께·수량의 블록 속성 태그 또는 문자 키워드", False),
    ("☐ 7_샘플기대값: 샘플 도면 DXF 15종 이상·STEP 5종 이상 + 사람이 잰 기대 메트릭", False),
    (
        "☐ 8_수기견적: 위 샘플 중 3건의 실제(수기) 견적서 금액 — 원 단위 일치 검증 기준(TC-67)",
        False,
    ),
    ("", False),
    ("사용처", True),
    (
        "1~6 시트 → 기준정보 화면 'G1 엑셀 가져오기'로 기준정보 초안에 반영(FN-16), 5·6은 2D 메트릭 규칙(FN-14).",
        False,
    ),
    (
        "7·8 시트 → 메트릭·원가 계산 회귀 테스트 정답지(TC-60, TC-67). 가져오기 대상이 아닙니다.",
        False,
    ),
]


def build_template(path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    head, example, inp = (PatternFill("solid", fgColor=c) for c in ("1F4E78", "EDEDED", "FFF2CC"))
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    wb = Workbook()
    ws = wb.active
    ws.title = "안내"
    for i, (text, bold) in enumerate(GUIDE, 1):
        ws.cell(row=i, column=1, value=text).font = Font(bold=bold, size=14 if i == 1 else 11)
    ws.column_dimensions["A"].width = 110
    for title, (desc, cols, examples) in SHEETS.items():
        ws = wb.create_sheet(title)
        ws["A1"], ws["A2"] = title, desc
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"].alignment = wrap
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(len(cols), 4))
        ws.row_dimensions[2].height = 48
        for c, (h, w, note) in enumerate(cols, 1):
            cell = ws.cell(row=HEADER_ROW, column=c, value=h)
            cell.font, cell.fill, cell.border, cell.alignment = (
                Font(bold=True, color="FFFFFF"),
                head,
                box,
                wrap,
            )
            ws.column_dimensions[cell.column_letter].width = w
            if note:
                cell.comment = Comment(note, "AX-CAD")
        blank = 0 if title == S_RATE else 20 if title == S_SAMPLE else 10
        for r, row in enumerate([*examples, *([()] * blank)], FIRST_ROW):
            for c in range(1, len(cols) + 1):
                cell = ws.cell(row=r, column=c, value=row[c - 1] if row else None)
                cell.fill = inp if not row or title == S_RATE else example
                cell.border, cell.alignment = box, wrap
        if title == S_RATE:  # the form rows: only the value column is input
            for r in range(FIRST_ROW, FIRST_ROW + len(examples)):
                for c in (1, 3, 4):
                    ws.cell(row=r, column=c).fill = example
        ws.freeze_panes = f"A{FIRST_ROW}"
    wb.save(path)


# --- import --------------------------------------------------------------------------------


class G1Error(Exception):
    """Whole-file problem (not a workbook, too big, missing sheet)."""


def _check_zip(path: str) -> None:
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
    except zipfile.BadZipFile:
        raise G1Error("엑셀(.xlsx) 파일이 아닙니다") from None
    if len(infos) > MAX_ZIP_ENTRIES or sum(i.file_size for i in infos) > MAX_UNZIPPED:
        raise G1Error("엑셀 파일 내용이 너무 큽니다")
    if "xl/workbook.xml" not in {i.filename for i in infos}:
        raise G1Error("엑셀(.xlsx) 파일이 아닙니다")


def _text(v: Any) -> str | None:
    s = str(v).strip() if v is not None else ""
    return s or None


def _num(v: Any) -> Decimal | None:
    s = _text(v)
    if s is None:
        return None
    try:
        d = Decimal(s.replace(",", "").removesuffix("%").strip())
    except InvalidOperation:
        raise ValueError(f"숫자가 아닙니다: {s[:40]}") from None
    if not d.is_finite():
        raise ValueError(f"숫자가 아닙니다: {s[:40]}")
    return d


def _pct(v: Any) -> str | None:
    d = _num(v)
    return None if d is None else str(d / 100)


def _rows(ws: Any, sheet: str) -> list[tuple[int, dict[str, Any]]]:
    headers = [h for h, _, _ in SHEETS[sheet][1]]
    got = [
        _text(c)
        for c in next(ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True), ())
    ]
    if got[: len(headers)] != headers:
        raise G1Error(f"{sheet} 시트의 4행 헤더가 양식과 다릅니다")
    out: list[tuple[int, dict[str, Any]]] = []
    for n, values in enumerate(
        ws.iter_rows(min_row=FIRST_ROW, max_col=len(headers), values_only=True), FIRST_ROW
    ):
        row = dict(zip(headers, values, strict=False))
        if all(_text(v) is None for v in values):
            continue
        if len(out) >= MAX_ROWS:
            raise G1Error(f"{sheet} 시트의 행이 {MAX_ROWS}개를 넘습니다")
        out.append((n, row))
    return out


_PARAM = re.compile(r"^\s*([A-Za-z0-9._\-]{1,40})\s*=\s*([^,]+?)\s*$")
TARGET = {"외곽선(절단)": "CUT", "절곡선": "BEND", "무시": "IGNORE", "표제란": "IGNORE"}
TITLE_FIELD = {
    "품번": "part_no",
    "품명": "part_name",
    "재질": "material",
    "두께": "thickness_mm",
    "수량": "qty",
}
RATE_KEYS = {
    "제조간접비율(%)": "overhead_rate",
    "일반관리비율(%)": "admin_rate",
    "이윤율(%)": "profit_rate",
    "VAT(%)": "vat_rate",
}


def _choice(v: Any, options: dict[str, str], what: str) -> str | None:
    s = _text(v)
    if s is None:
        return None
    for key, code in options.items():
        if key in s:
            return code
    raise ValueError(f"{what}: '{s[:40]}'은(는) 선택지({' / '.join(options)})가 아닙니다")


def parse_g1(path: str) -> dict[str, Any]:
    """-> {bundle, warnings, errors}. Row problems are collected (sheet + row number), not raised,
    so the admin sees every mistake at once."""
    from openpyxl import load_workbook

    _check_zip(path)
    wb = load_workbook(
        path, read_only=True, data_only=True
    )  # data_only: cached values, no formulas
    missing = [
        s for s in (S_MAT, S_PROC, S_LAB, S_RATE, S_LAYER, S_TITLE) if s not in wb.sheetnames
    ]
    if missing:
        raise G1Error("시트가 없습니다: " + ", ".join(missing))
    errors: list[str] = []
    skipped = 0
    bundle: dict[str, Any] = {
        "materials": [],
        "price_items": [],
        "process_rules": [],
        "mapping_rules": [],
        "cost_ratios": None,
    }

    def rows(sheet: str) -> list[tuple[int, dict[str, Any]]]:
        nonlocal skipped
        out: list[tuple[int, dict[str, Any]]] = []
        for n, r in _rows(wb[sheet], sheet):
            if (_text(r.get("비고")) or "").startswith(EXAMPLE) and sheet != S_RATE:
                skipped += 1
                continue
            out.append((n, r))
        return out

    def guard(sheet: str, n: int, fn: Any) -> None:
        try:
            fn()
        except (ValueError, KeyError) as e:
            errors.append(f"{sheet} {n}행: {e}")

    for n, r in rows(S_MAT):

        def mat(r: dict[str, Any] = r) -> None:
            lo = hi = None
            span = _text(r["두께 범위(mm)"])
            if span and span != "전체":
                a, _, b = span.replace("~", "-").partition("-")
                lo, hi = _num(a), _num(b or a)
            bundle["materials"].append(
                {
                    "material_code": _text(r["재질코드"]),
                    "material_name": _text(r["재질명"]),
                    "thickness_min_mm": None if lo is None else str(lo),
                    "thickness_max_mm": None if hi is None else str(hi),
                    "density_g_cm3": str(_num(r["비중(g/cm³)"])),
                    "unit_price_per_kg": None if (p := _num(r["단가(원/kg)"])) is None else str(p),
                    "scrap_rate": _pct(r["스크랩률(%)"]),
                }
            )

        guard(S_MAT, n, mat)

    labor = {}
    for n, r in rows(S_LAB):

        def lab(r: dict[str, Any] = r) -> None:
            code = _text(r["구분"])
            if not code:
                raise ValueError("구분이 비어 있습니다")
            price = _num(r["임률(원/h)"])
            labor[code] = True
            bundle["price_items"].append(
                {
                    "item_code": code,
                    "item_type": "LABOR",
                    "unit": "h",
                    "unit_price": None if price is None else str(price),
                }
            )

        guard(S_LAB, n, lab)

    for n, r in rows(S_PROC):

        def proc(r: dict[str, Any] = r) -> None:
            code = _text(r["공정코드"])
            params = {}
            for part in (_text(r["파라미터 값"]) or "").split(","):
                if not part.strip():
                    continue
                m = _PARAM.match(part)
                if not m:
                    raise ValueError(f"파라미터 형식은 이름=숫자 입니다: '{part.strip()[:40]}'")
                params[m.group(1)] = float(_num(m.group(2)) or 0)
            machine = None
            if (cost := _num(r["기계경비(원/h)"])) is not None:
                machine = f"{code}-MC"
                bundle["price_items"].append(
                    {
                        "item_code": machine,
                        "item_type": "MACHINE",
                        "unit": "h",
                        "unit_price": str(cost),
                    }
                )
            lab_code = _text(r["임률 구분"])
            if lab_code and lab_code not in labor:
                raise ValueError(f"임률 구분 '{lab_code}'이(가) 3_임률 시트에 없습니다")
            bundle["process_rules"].append(
                {
                    "rule_code": code,
                    "process_code": code,
                    "process_name": _text(r["공정명"]),
                    "input_metric": _text(r["입력 메트릭"]),
                    "formula_text": _text(r["표준공수 공식(분)"]),
                    "params": params,
                    "labor_item_code": lab_code,
                    "machine_item_code": machine,
                }
            )

        guard(S_PROC, n, proc)

    ratios: dict[str, Any] = {}
    for n, r in rows(S_RATE):

        def rate(r: dict[str, Any] = r) -> None:
            item, v = _text(r["항목"]) or "", r["값"]
            if item in RATE_KEYS:
                ratios[RATE_KEYS[item]] = _pct(v)
            elif item == "제조간접비 배부 기준":
                ratios["overhead_basis"] = _choice(
                    v, {"기계": "MACHINE_HOUR", "노무비": "LABOR_RATIO"}, item
                )
            elif item == "절사 규칙":
                ratios["rounding_rule"] = _choice(
                    v, {"절사": "FLOOR", "반올림": "HALF_UP", "올림": "CEILING"}, item
                )
            elif item == "절사 단위(원)":
                u = _num(v)
                ratios["rounding_unit"] = None if u is None else int(u)
            elif item == "절사 적용 시점":
                ratios["rounding_scope"] = _choice(v, {"라인": "LINE", "합계": "TOTAL"}, item)

        guard(S_RATE, n, rate)
    required = ("overhead_basis", "rounding_rule", "rounding_unit", "rounding_scope")
    if any(ratios.get(k) is not None for k in (*required, "admin_rate", "profit_rate")):
        empty = [k for k in required if ratios.get(k) is None]
        if empty:
            errors.append(f"{S_RATE}: 함께 입력해야 하는 값이 비어 있습니다 ({', '.join(empty)})")
        else:
            bundle["cost_ratios"] = {k: v for k, v in ratios.items() if v is not None}

    for n, r in rows(S_LAYER):

        def layer(r: dict[str, Any] = r) -> None:
            use = _text(r["용도"]) or ""
            if use.startswith("펀칭"):
                dia = _num(
                    r["레이어명 패턴"].replace("≤", "")
                    if isinstance(r["레이어명 패턴"], str)
                    else r["레이어명 패턴"]
                )
                if dia is None:
                    raise ValueError("펀칭 기준 직경(mm)을 숫자로 적어 주세요")
                bundle["mapping_rules"].append(
                    {"rule_type": "PUNCH_MAX_DIA", "target": "HOLE", "pattern": str(dia)}
                )
                return
            if use not in TARGET:
                raise ValueError(
                    f"용도 '{use[:40]}'은(는) {' / '.join(TARGET)} / 펀칭 기준 직경(mm) 중 하나여야 합니다"
                )
            for kind, col in (("LAYER", "레이어명 패턴"), ("LINETYPE", "선종(linetype)")):
                if pat := _text(r[col]):
                    bundle["mapping_rules"].append(
                        {"rule_type": kind, "target": TARGET[use], "pattern": pat}
                    )

        guard(S_LAYER, n, layer)

    for n, r in rows(S_TITLE):

        def title(r: dict[str, Any] = r) -> None:
            item = _text(r["항목"]) or ""
            if item not in TITLE_FIELD:
                raise ValueError(
                    f"항목 '{item[:40]}'은(는) {' / '.join(TITLE_FIELD)} 중 하나여야 합니다"
                )
            pats = [p for p in (_text(r["블록 속성 태그"]), _text(r["문자 키워드"])) if p]
            if pats:
                bundle["mapping_rules"].append(
                    {
                        "rule_type": "TITLE_TAG",
                        "target": TITLE_FIELD[item],
                        "pattern": ", ".join(pats),
                    }
                )

        guard(S_TITLE, n, title)

    warnings = [f"예시 행 {skipped}개를 건너뛰었습니다"] if skipped else []
    return {"bundle": bundle, "warnings": warnings, "errors": errors}


def parse_g1_job(path: str) -> str:
    """Worker entry (run_isolated): returns JSON text, G1Error becomes an `errors` entry."""
    import json

    try:
        return json.dumps(parse_g1(path), ensure_ascii=False)
    except G1Error as e:
        return json.dumps({"bundle": None, "warnings": [], "errors": [str(e)]}, ensure_ascii=False)
